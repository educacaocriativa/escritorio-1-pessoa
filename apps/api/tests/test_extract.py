"""Story 7.7 — caminho feliz estendido + caminho infeliz de core/extract.py.

Exercita `extract_text` além do `.txt` (PDF/docx/imagem-OCR) e os caminhos infelizes
(tipo não suportado → marcador explícito; arquivo corrompido → marcador de erro). Gera os
arquivos reais EM MEMÓRIA via as libs já dependências (fpdf2/python-docx/Pillow) — sem
fixtures binárias versionadas. O teste de OCR pula de forma EXPLÍCITA (skipif) quando o
binário `tesseract` não está no ambiente (ver CLAUDE.md §6.1); roda de verdade na imagem
Docker de produção/CI, que instala tesseract-ocr/tesseract-ocr-por.
"""
import io
import shutil

import pytest

from app.core.extract import extract_text

_DOCX_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


# --- AC1: caminho feliz estendido ------------------------------------------------


def test_extract_text_pdf_happy_path() -> None:
    from fpdf import FPDF

    known = "Peticao inicial de teste"
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "", 14)
    pdf.cell(0, 10, known)
    data = bytes(pdf.output())

    result = extract_text(data, "application/pdf", "peticao.pdf")

    assert known in result


def test_extract_text_docx_happy_path() -> None:
    from docx import Document

    known = "Contrato de prestacao de servicos de teste"
    doc = Document()
    doc.add_paragraph(known)
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()

    result = extract_text(data, _DOCX_CT, "contrato.docx")

    assert known in result


@pytest.mark.skipif(
    shutil.which("tesseract") is None,
    reason=(
        "tesseract não instalado neste ambiente — ver CLAUDE.md §6.1; presente na imagem "
        "Docker de produção via apps/api/Dockerfile linhas 10-12"
    ),
)
def test_extract_text_image_ocr_happy_path() -> None:
    from PIL import Image, ImageDraw, ImageFont

    expected = "TESTE"
    img = Image.new("RGB", (400, 120), color="white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 72)
    except OSError:
        font = ImageFont.load_default()
    draw.text((20, 20), expected, fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    data = buf.getvalue()

    result = extract_text(data, "image/png", "documento.png")

    # OCR não é 100% determinístico — tolerante (case-insensitive, substring).
    assert expected.lower() in result.lower()


# --- AC2: caminho infeliz --------------------------------------------------------


def test_extract_text_unsupported_type_is_explicit_not_silent_empty() -> None:
    # application/zip não casa nenhuma branch de content_type/extensão.
    result = extract_text(b"conteudo qualquer de um arquivo zip", "application/zip", "arquivo.zip")

    assert result != ""
    assert "não suportado" in result.lower() or "nao suportado" in result.lower()


def test_extract_text_corrupted_pdf_returns_explicit_error_marker() -> None:
    result = extract_text(b"isto nao e um pdf valido", "application/pdf", "corrompido.pdf")

    assert result.startswith("[ERRO NA EXTRAÇÃO:")


def test_extract_text_corrupted_docx_returns_explicit_error_marker() -> None:
    result = extract_text(b"isto nao e um docx valido", _DOCX_CT, "corrompido.docx")

    assert result.startswith("[ERRO NA EXTRAÇÃO:")
